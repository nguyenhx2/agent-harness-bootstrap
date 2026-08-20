#!/usr/bin/env python3
"""Decide, per source document, which reader owns its text - and say so when none can.

The mechanism is distilled from the docs-to-knowledge skill (its scripts/route_sources.py
and references/format_routing.md): classify, probe, choose a strategy, carry a reason string
for every decision. Left behind is that skill's output contract - the knowledge base, the
render and audit layers - because spec-builder's output is its numbered sections.

The one decision this exists for: **knowing when a source cannot be read.** A scanned PDF
still has a text layer; it is just a lie. Extracting it "successfully" returns a handful of
characters per page, and a model handed near-empty text about a document the user called the
requirements fills the gap with something plausible. That is the invented requirement the
never-invent rule exists to prevent, arriving through the file reader instead of through the
interview. Below TEXT_LAYER_MIN_CHARS_PER_PAGE the route switches to vision and says why.

Optional dependencies degrade HONESTLY: a missing reader changes the route and names what
was lost, never crashes and never silently skips a file. A source nothing installed can read
routes to `unreadable`, and per reference/elicitation.md that becomes an OI-nn in section 11
naming the file and the fix - never a silent gap.

Usage:
    python route_sources.py <path-or-url> [more...] [-o plan.json] [--out-dir DIR]
    python route_sources.py ./inputs --recursive -o plan.json
    python route_sources.py deck.pptx --prefer-markitdown     # override the policy
    python route_sources.py --self-test                       # policy table only
"""

from __future__ import annotations

import argparse
import importlib
import json
import os
import re
import shutil
import sys

HERE = os.path.dirname(os.path.abspath(__file__))

# A PDF page carrying fewer than this many extracted characters has no usable text layer:
# it is a scan, or a deck exported as flat images. See the module docstring.
TEXT_LAYER_MIN_CHARS_PER_PAGE = 80

NATIVE_DOC = {".pdf": "pdf", ".docx": "docx", ".pptx": "pptx",
              ".xlsx": "xlsx", ".xlsm": "xlsx"}
LEGACY_DOC = {".doc": "doc", ".ppt": "ppt", ".xls": "xls", ".rtf": "rtf",
              ".odt": "odt", ".ods": "ods", ".odp": "odp", ".xlsb": "xlsb"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".tif", ".tiff"}
MARKITDOWN_EXT = {".html": "html", ".htm": "html", ".csv": "csv", ".tsv": "csv",
                  ".json": "json", ".xml": "xml", ".msg": "msg", ".eml": "msg",
                  ".epub": "epub", ".ipynb": "ipynb", ".zip": "zip"}
PLAIN_EXT = {".md": "markdown", ".markdown": "markdown", ".txt": "text", ".log": "text"}
SKIP_NAMES = ("~$", ".ds_store", "thumbs.db")


# ------------------------------------------------------------------ capabilities
def have(mod: str) -> bool:
    try:
        importlib.import_module(mod)
        return True
    except Exception:
        return False


def have_markitdown() -> bool:
    return bool(shutil.which("markitdown")) or have("markitdown")


def have_anydoc() -> bool:
    return have("anydoc")  # pip install firecrawl-anydoc


def detect_caps() -> dict[str, bool]:
    """What this machine can actually read. Passed into the policy rather than probed
    inside it, so the self-test can exercise every branch without installing anything."""
    return {"pymupdf": have("fitz"), "docx": have("docx"), "pptx": have("pptx"),
            "openpyxl": have("openpyxl"), "anydoc": have_anydoc(),
            "markitdown": have_markitdown()}


# ------------------------------------------------------------------ classify
def classify(path: str) -> tuple[str, str]:
    """-> (format id, family) where family is office | legacy | image | markitdown |
    plain | url | unknown."""
    if re.match(r"^https?://", path, re.I):
        return ("url", "url")
    ext = os.path.splitext(path)[1].lower()
    if ext in NATIVE_DOC:
        return (NATIVE_DOC[ext], "office")
    if ext in LEGACY_DOC:
        return (LEGACY_DOC[ext], "legacy")
    if ext in IMAGE_EXT:
        return ("image", "image")
    if ext in MARKITDOWN_EXT:
        return (MARKITDOWN_EXT[ext], "markitdown")
    if ext in PLAIN_EXT:
        return (PLAIN_EXT[ext], "plain")
    return (ext.lstrip(".") or "noext", "unknown")


def probe(path: str, fmt: str, family: str, caps: dict[str, bool]) -> dict:
    """Cheap inspection that informs the route. Never fatal: a locked or corrupt file must
    still appear in the plan, with its error recorded."""
    info: dict = {}
    if family == "url":
        return info
    try:
        info["size_kb"] = round(os.path.getsize(path) / 1024, 1)
    except OSError:
        pass
    try:
        if fmt == "pdf" and caps["pymupdf"]:
            import fitz
            doc = fitz.open(path)
            chars = sum(len(p.get_text().strip()) for p in doc)
            info["pages"] = doc.page_count
            info["chars_per_page"] = round(chars / max(doc.page_count, 1), 1)
            info["has_text_layer"] = (info["chars_per_page"]
                                      >= TEXT_LAYER_MIN_CHARS_PER_PAGE)
            doc.close()
        elif fmt == "pptx" and caps["pptx"]:
            from pptx import Presentation
            info["slides"] = len(Presentation(path).slides)
        elif fmt == "xlsx" and caps["openpyxl"]:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            info["sheets"] = wb.sheetnames
            wb.close()
    except Exception as e:  # a bad file is a finding, not a crash
        info["probe_error"] = f"{type(e).__name__}: {e}"
    return info


# ------------------------------------------------------------------ POLICY
def choose_text_strategy(fmt: str, family: str, info: dict, caps: dict[str, bool],
                         prefer: str = "auto") -> tuple[str, str]:
    """POLICY HOOK - who owns the TEXT layer. -> (strategy, reason).

    native      the format has an object model; read it directly, keep page/sheet/slide
                boundaries
    anydoc      no native reader, or Markdown tables beat cell coordinates (spreadsheets)
    markitdown  no native reader here and no anydoc; broad coverage, flat structure
    vision      no trustworthy text layer exists; the pixels are the source
    read        already text
    unreadable  nothing installed can read it. Not a failure to hide - an OI-nn.
    """
    def fallback(why: str, install: str) -> tuple[str, str]:
        if caps["anydoc"]:
            return ("anydoc", f"{why}; anydoc reads it directly")
        if caps["markitdown"]:
            return ("markitdown", f"{why}; markitdown converts it, losing page boundaries")
        return ("unreadable", f"{why}, and no converter is installed ({install})")

    if prefer == "markitdown" and family in ("office", "legacy", "markitdown", "plain",
                                             "url"):
        if caps["markitdown"]:
            return ("markitdown", "forced by --prefer-markitdown")
        return ("unreadable", "--prefer-markitdown, but markitdown is not installed "
                              "(pip install 'markitdown[all]')")
    if prefer == "native" and family == "office":
        return ("native", "forced by --prefer-native")

    if family == "url":
        if caps["markitdown"]:
            return ("markitdown", "remote content; markitdown fetches and converts")
        return ("unreadable", "remote content and markitdown is not installed "
                              "(pip install 'markitdown[all]')")
    if family == "plain":
        return ("read", "already text; read it directly")
    if family == "image":
        return ("vision", "pixels only; a vision agent must read it")

    if fmt == "pdf":
        if info.get("has_text_layer") is False:
            return ("vision", f"scanned or flat PDF ({info.get('chars_per_page', 0)} "
                              f"chars/page, under {TEXT_LAYER_MIN_CHARS_PER_PAGE}); the "
                              f"text layer is a lie, the pixels are the source")
        if caps["pymupdf"]:
            return ("native", "PDF with a real text layer; pymupdf keeps page boundaries")
        return fallback("pymupdf is not installed, so the text layer could not even be "
                        "probed for a scan (pip install pymupdf reads PDF natively)",
                        "pip install pymupdf")
    if fmt == "docx":
        if caps["docx"]:
            return ("native", "paragraph and table structure is worth keeping")
        return fallback("python-docx is not installed", "pip install python-docx")
    if fmt == "pptx":
        if caps["pptx"]:
            return ("native", "shape walk keeps slide boundaries and table cells")
        return fallback("python-pptx is not installed", "pip install python-pptx")
    if fmt == "xlsx":
        # A requirements spreadsheet is read as tables, not as coordinates: anydoc and
        # markitdown emit real Markdown tables, openpyxl emits tab-joined rows.
        if caps["anydoc"]:
            return ("anydoc", "spreadsheets read better as Markdown tables; anydoc is "
                              "the fastest local reader")
        if caps["markitdown"]:
            return ("markitdown", "spreadsheets read better as Markdown tables")
        if caps["openpyxl"]:
            return ("native", "openpyxl gives tab-joined rows per sheet, not Markdown "
                              "tables (pip install firecrawl-anydoc for real tables)")
        return ("unreadable", "no spreadsheet reader installed (pip install openpyxl, "
                              "or firecrawl-anydoc for Markdown tables)")
    if family == "legacy":
        return fallback(f"no native reader for .{fmt} (legacy binary or OpenDocument)",
                        "pip install firecrawl-anydoc")
    if family == "markitdown":
        if caps["markitdown"]:
            return ("markitdown", f"no native reader for .{fmt} in this skill")
        return ("unreadable", f"no native reader for .{fmt} and markitdown is not "
                              f"installed (pip install 'markitdown[all]')")
    if caps["markitdown"]:
        return ("markitdown", "unknown extension; try markitdown, then report the result")
    return ("unreadable", "unknown extension and no converter installed; ask the user to "
                          "re-export as PDF, DOCX or Markdown")


def choose_visual_strategy(family: str, text_strategy: str) -> tuple[str, str]:
    """Rendering is narrow here on purpose: pixels are produced only when the text is not
    trustworthy. docs-to-knowledge rasterizes everything because it cross-checks diagrams
    against text; a BA reading a readable document gains nothing from a PNG of it."""
    if family == "image":
        return ("copy", "the source already is an image")
    if text_strategy == "vision":
        return ("render", "no trustworthy text layer; the pages ARE the source")
    return ("none", "the text layer is readable; rasterizing it adds nothing")


# ------------------------------------------------------------------ self-test
# A policy table that cannot be exercised without installing six optional packages is a
# policy nobody tests. Capabilities are injected, so every branch runs everywhere.
SELF_TEST_CASES = [
    # (label, fmt, family, probe info, caps overrides, expected strategy)
    ("scanned pdf", "pdf", "office", {"has_text_layer": False, "chars_per_page": 11.0},
     {"pymupdf": True}, "vision"),
    ("pdf with a text layer", "pdf", "office", {"has_text_layer": True},
     {"pymupdf": True}, "native"),
    ("pdf, no pymupdf", "pdf", "office", {}, {"markitdown": True}, "markitdown"),
    ("xlsx, anydoc present", "xlsx", "office", {}, {"anydoc": True}, "anydoc"),
    ("xlsx, anydoc absent", "xlsx", "office", {}, {"markitdown": True}, "markitdown"),
    ("xlsx, nothing installed", "xlsx", "office", {}, {}, "unreadable"),
    ("docx", "docx", "office", {}, {"docx": True}, "native"),
    ("pptx", "pptx", "office", {}, {"pptx": True}, "native"),
    ("legacy .doc, anydoc present", "doc", "legacy", {}, {"anydoc": True}, "anydoc"),
    ("markdown", "markdown", "plain", {}, {}, "read"),
    ("image", "image", "image", {}, {}, "vision"),
    ("url", "url", "url", {}, {"markitdown": True}, "markitdown"),
    ("unknown ext", "wat", "unknown", {}, {"markitdown": True}, "markitdown"),
    ("unknown ext, nothing installed", "wat", "unknown", {}, {}, "unreadable"),
]
NO_CAPS = {"pymupdf": False, "docx": False, "pptx": False, "openpyxl": False,
           "anydoc": False, "markitdown": False}


def self_test() -> int:
    problems = []
    for label, fmt, family, info, caps_on, expect in SELF_TEST_CASES:
        caps = dict(NO_CAPS, **caps_on)
        got, reason = choose_text_strategy(fmt, family, info, caps)
        ok = got == expect
        print(f"    {'ok  ' if ok else 'FAIL'} {label:<32} -> {got:<11} {reason}")
        if not ok:
            problems.append(f"{label}: expected {expect}, got {got}")
        if not reason.strip():
            problems.append(f"{label}: routed with no reason string")

    # DEAD CHECK 1. The 80-chars/page rule is the whole point. If the same PDF routes the
    # same way whether or not it has a text layer, the probe is decorative.
    caps = dict(NO_CAPS, pymupdf=True)
    scanned = choose_text_strategy("pdf", "office", {"has_text_layer": False}, caps)[0]
    real = choose_text_strategy("pdf", "office", {"has_text_layer": True}, caps)[0]
    if scanned == real:
        problems.append("DEAD CHECK: the text-layer probe changes nothing, so a scanned "
                        "PDF would be read as though its text were real")

    # DEAD CHECK 2. Capabilities are injected so the table can run anywhere; if injecting
    # them changes no answer, the injection is not wired and this table proves nothing.
    if (choose_text_strategy("xlsx", "office", {}, dict(NO_CAPS, anydoc=True))[0]
            == choose_text_strategy("xlsx", "office", {}, dict(NO_CAPS))[0]):
        problems.append("DEAD CHECK: capability flags do not reach the policy, so every "
                        "case above is testing one hard-coded answer")

    # DEAD CHECK 3. `unreadable` is what becomes an OI-nn. A policy that always finds some
    # route hides the unreadable source instead of reporting it.
    if all(choose_text_strategy(f, fam, {}, dict(NO_CAPS))[0] != "unreadable"
           for f, fam in (("xlsx", "office"), ("doc", "legacy"), ("wat", "unknown"))):
        problems.append("DEAD CHECK: nothing routes to `unreadable`, so a source no "
                        "installed reader can open would be silently skipped")

    if problems:
        print("\n  routing policy is broken:")
        for p in problems:
            print(f"    {p}")
        return 1
    print(f"    self-test: {len(SELF_TEST_CASES)} routing cases, and the text-layer probe, "
          f"the capability flags and the unreadable route all still change the answer.")
    return 0


# ------------------------------------------------------------------ plan
def slugify(name: str, taken: set[str]) -> str:
    base = os.path.splitext(os.path.basename(name))[0]
    base = re.sub(r"[^0-9A-Za-z]+", "_", base).strip("_").lower() or "src"
    slug, n = base, 2
    while slug in taken:
        slug, n = f"{base}_{n}", n + 1
    taken.add(slug)
    return slug


def build_entry(path: str, out_dir: str, taken: set[str], caps: dict[str, bool],
                prefer: str) -> dict:
    fmt, family = classify(path)
    info = probe(path, fmt, family, caps)
    tstrat, treason = choose_text_strategy(fmt, family, info, caps, prefer)
    vstrat, vreason = choose_visual_strategy(family, tstrat)
    slug = slugify(path if family != "url" else re.sub(r"\W+", "_", path)[:40], taken)
    ext = ".md" if tstrat in ("anydoc", "markitdown") else ".txt"
    out = os.path.join(out_dir, slug + ext)

    steps, notes = [], []
    if tstrat == "unreadable":
        notes.append("UNREADABLE: record an OI-nn in section 11 naming this file, why it "
                     "cannot be read, and the fix - the install above, or a re-export "
                     "request to the user. Never a silent gap.")
    elif tstrat == "vision":
        notes.append("No text to extract. Read the rendered pages (or the image itself) "
                     "with a vision agent, and say in the spec that the content was read "
                     "from pixels.")
        steps.append(f'"{sys.executable}" "{os.path.join(HERE, "ingest.py")}" "{path}" '
                     f'--via vision --render "{os.path.join(out_dir, slug + "_pages")}"')
    else:
        steps.append(f'"{sys.executable}" "{os.path.join(HERE, "ingest.py")}" "{path}" '
                     f'--via {tstrat} --out "{out}"')
    if info.get("probe_error"):
        notes.append(f"Probe failed ({info['probe_error']}): the route below is a guess "
                     "from the extension alone. If ingest also fails, that is an OI-nn.")
    return {"path": path, "name": os.path.basename(path) or path, "slug": slug,
            "format": fmt, "family": family, "text_strategy": tstrat,
            "text_reason": treason, "visual_strategy": vstrat, "visual_reason": vreason,
            "out": out if tstrat not in ("unreadable", "vision") else None,
            "steps": steps, "probe": info, "notes": notes}


def expand(paths: list[str], recursive: bool) -> list[str]:
    out = []
    for p in paths:
        if re.match(r"^https?://", p, re.I):
            out.append(p)
        elif os.path.isdir(p):
            walker = os.walk(p) if recursive else [(p, [], sorted(os.listdir(p)))]
            for root, _dirs, files in walker:
                for f in sorted(files):
                    low = f.lower()
                    if low.startswith(SKIP_NAMES) or low in SKIP_NAMES:
                        continue
                    out.append(os.path.join(root, f))
        elif os.path.exists(p):
            out.append(p)
        else:
            sys.stderr.write(f"[warn] not found, skipped: {p}\n")
    return out


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):  # CJK/VI filenames vs a cp1252 console
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Route source documents to their readers.")
    ap.add_argument("sources", nargs="*", help="files, folders and/or URLs")
    ap.add_argument("-o", "--out", help="write the plan JSON here")
    ap.add_argument("--out-dir", default="sources_text",
                    help="where ingest.py writes the extracted text")
    ap.add_argument("--recursive", action="store_true", default=True)
    ap.add_argument("--no-recursive", dest="recursive", action="store_false")
    ap.add_argument("--prefer-markitdown", action="store_const", const="markitdown",
                    dest="prefer", help="override: markitdown owns every text layer")
    ap.add_argument("--prefer-native", action="store_const", const="native", dest="prefer")
    ap.add_argument("--self-test", action="store_true",
                    help="run the routing policy table and exit")
    ap.set_defaults(prefer="auto")
    a = ap.parse_args()

    # The policy table runs on every invocation, not only under --self-test: it is
    # dependency-free and costs milliseconds, and a routing bug found after ingestion has
    # already written twelve files is found too late.
    print("  routing policy:")
    if self_test() != 0:
        return 1
    if a.self_test:
        return 0
    if not a.sources:
        sys.stderr.write("\nNo sources given. Pass files, folders or URLs.\n")
        return 2

    caps = detect_caps()
    print("\n  readers installed: " + ", ".join(
        f"{'ok' if v else 'MISSING'} {k}" for k, v in sorted(caps.items())))

    files = expand(a.sources, a.recursive)
    if not files:
        sys.stderr.write("No readable sources.\n")
        return 2
    taken: set[str] = set()
    entries = [build_entry(f, a.out_dir, taken, caps, a.prefer) for f in files]

    w = max([len(e["name"]) for e in entries] + [6])
    print(f"\n  {'SOURCE'.ljust(w)}  {'FORMAT':<8} {'TEXT':<11} {'VISUAL':<7} WHY")
    print("  " + "-" * (w + 42))
    for e in entries:
        print(f"  {e['name'].ljust(w)}  {e['format']:<8} {e['text_strategy']:<11} "
              f"{e['visual_strategy']:<7} {e['text_reason']}")

    blocked = [e for e in entries if e["text_strategy"] == "unreadable"]
    if blocked:
        print("\n  UNREADABLE - one OI-nn each in section 11, naming the file and the fix:")
        for e in blocked:
            print(f"    {e['name']}: {e['text_reason']}")
    if any(e["text_strategy"] == "vision" for e in entries):
        print("\n  At least one source has no trustworthy text layer: its pages must be "
              "read by a vision agent, and the spec must say so.")

    plan = {"version": 1, "out_dir": a.out_dir, "n_sources": len(entries),
            "capabilities": caps, "sources": entries}
    if a.out:
        os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
        with open(a.out, "w", encoding="utf-8", newline="\n") as f:
            json.dump(plan, f, ensure_ascii=False, indent=1)
        print(f"\n  plan -> {a.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
